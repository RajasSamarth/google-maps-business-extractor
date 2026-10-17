"""Excel export helpers for the business extractor."""
from pathlib import Path
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from app.schemas import BusinessRecord, SearchInfo


# Exact column names and order required by the project/problem statement,
# plus a trailing status column showing skipped detail tabs.
COLUMNS = [
    "Sr. No.",
    "Business Name",
    "Category",
    "Full Address",
    "Area / Locality",
    "City / Pincode",
    "Phone Number",
    "Website",
    "Rating",
    "Total Reviews",
    "Opening Hours",
    "Business Status",
    "Latitude, Longitude",
    "Google Maps Link",
    "Place ID",
    "Detail Status",
]


def _as_dict(record) -> dict:
    """Accept a BusinessRecord or a plain dict and return a dict."""
    if isinstance(record, BusinessRecord):
        return record.to_dict()
    return dict(record or {})


def _clean(value):
    """Convert empty strings / NaN-like values to None (blank cell)."""
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped != "" else None
    # pandas NaN guard (optional dependency not required here)
    try:
        import math
        if isinstance(value, float) and math.isnan(value):
            return None
    except Exception:
        pass
    return value


def sanitize_filename_part(value: str) -> str:
    """Make a string safe to use as part of a file name."""
    value = (value or "").strip().lower()
    # Replace anything that is not a letter, digit, or underscore with underscore.
    value = re.sub(r"[^a-z0-9]+", "_", value)
    value = re.sub(r"_+", "_", value).strip("_")
    return value or "search"


def build_filename(query: str, location: str, when: date | None = None) -> str:
    q = sanitize_filename_part(query)[:30]
    loc = sanitize_filename_part(location)[:60]
    day = (when or date.today()).isoformat()
    return f"{q}_{loc}_{day}.xlsx"


def _coords(record_dict: dict) -> str | None:
    lat = _clean(record_dict.get("latitude"))
    lon = _clean(record_dict.get("longitude"))
    if lat is None and lon is None:
        return None
    if lat is None:
        return str(lon)
    if lon is None:
        return str(lat)
    return f"{lat}, {lon}"


def _detail_status(record_dict: dict) -> str:
    return "Skipped" if record_dict.get("detail_failed") else "Complete"


def _row_from_record(index: int, record) -> list:
    d = _as_dict(record)
    return [
        index,
        _clean(d.get("business_name")),
        _clean(d.get("category")),
        _clean(d.get("full_address")),
        _clean(d.get("area_locality")),
        _clean(d.get("city_pincode")),
        _clean(d.get("phone_number")),
        _clean(d.get("website")),
        _clean(d.get("rating")),
        _clean(d.get("total_reviews")),
        _clean(d.get("opening_hours")),
        _clean(d.get("business_status")),
        _coords(d),
        _clean(d.get("google_maps_link")),
        _clean(d.get("place_id")),
        _detail_status(d),
    ]


def export_to_excel(records, search_info, output_dir="output", filename_location=None) -> Path:
    """Create an .xlsx file from records and search metadata.

    ``filename_location`` optionally overrides the location text used only
    for the file name (multi-area searches would otherwise produce
    over-long file names). Returns the path to the generated file.
    """
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)

    if isinstance(search_info, SearchInfo):
        info = search_info
    else:
        from app.schemas import SearchInfo as _SI
        info = _SI(
            query=search_info.get("query", ""),
            location=search_info.get("location", ""),
            radius_km=search_info.get("radius_km"),
            timestamp=search_info.get("timestamp", "") or None,
            total_results=search_info.get("total_results", len(list(records))),
        )

    records_list = list(records or [])
    name_query = info.query
    name_location = filename_location if filename_location else info.location
    filename = build_filename(name_query, name_location)
    file_path = output_path / filename

    wb = Workbook()

    # --- First worksheet: business records ---
    ws = wb.active
    ws.title = "Businesses"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, rec in enumerate(records_list, start=1):
        ws.append(_row_from_record(i, rec))

    # Light column-width tuning (optional, keeps sheet readable)
    widths = [8, 28, 20, 40, 22, 18, 16, 30, 8, 10, 26, 16, 22, 40, 24, 14]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    ws.freeze_panes = "A2"

    # --- Second worksheet: search info ---
    ws_info = wb.create_sheet("Search Info")
    ws_info.append(["Field", "Value"])
    for cell in ws_info[1]:
        cell.font = Font(bold=True)
    info_rows = [
        ("Search query", info.query),
        ("Location", info.location),
        ("Search radius", (f"{info.radius_km} km" if info.radius_km is not None else "")),
        ("Search timestamp", info.timestamp),
        ("Total result count", info.total_results),
    ]
    for key, value in info_rows:
        ws_info.append([key, value])
    ws_info.column_dimensions["A"].width = 22
    ws_info.column_dimensions["B"].width = 40

    wb.save(file_path)
    return file_path
