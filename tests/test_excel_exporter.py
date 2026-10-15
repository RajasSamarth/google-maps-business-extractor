"""Tests for Excel exporter using synthetic data."""
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from app.schemas import BusinessRecord, SearchInfo
from app.services.excel_exporter import (
    COLUMNS,
    build_filename,
    export_to_excel,
    sanitize_filename_part,
)


def _sample_records():
    return [
        BusinessRecord(
            business_name="Dharampeth Dental Care",
            category="Dental clinic",
            full_address="123, Wardha Road, Dharampeth",
            area_locality="Dharampeth",
            city_pincode="Nagpur 440010",
            phone_number="+91 9876543210",
            website="https://example.com",
            rating=4.6,
            total_reviews=120,
            opening_hours="Mon-Sat 10:00-19:00",
            business_status="Operational",
            latitude=21.1458,
            longitude=79.0882,
            google_maps_link="https://maps.example/1",
            place_id="place_1",
        ),
        # Record with missing fields to verify blanks.
        BusinessRecord(business_name="City Smile Clinic", rating=4.2, place_id="place_2"),
        # Duplicate Place ID case is not removed here (cleaner responsibility);
        # exporter just writes what it receives.
        BusinessRecord(business_name="Third Clinic", place_id="place_1"),
    ]


def test_export_creates_two_sheets_and_columns(tmp_path):
    out_dir = tmp_path / "output"
    records = _sample_records()
    info = SearchInfo(
        query="dental clinic", location="Dharampeth, Nagpur", radius_km=3, total_results=len(records)
    )
    file_path = export_to_excel(records, info, output_dir=out_dir)

    assert file_path.exists()
    assert file_path.name == build_filename("dental clinic", "Dharampeth, Nagpur")

    wb = load_workbook(file_path)
    assert set(wb.sheetnames) == {"Businesses", "Search Info"}

    ws = wb["Businesses"]
    headers = [c.value for c in ws[1]]
    assert headers == COLUMNS
    # 3 records + header
    assert ws.max_row == 4

    # Check first data row values.
    row2 = [c.value for c in ws[2]]
    assert row2[1] == "Dharampeth Dental Care"
    assert row2[8] == 4.6
    assert row2[12] == "21.1458, 79.0882"

    # Second record has blanks for most fields.
    row3 = [c.value for c in ws[3]]
    assert row3[1] == "City Smile Clinic"
    assert row3[2] is None  # category blank
    assert row3[12] is None  # coords blank
    assert row3[14] == "place_2"

    # Search Info sheet
    si = wb["Search Info"]
    info_map = {r[0]: r[1] for r in si.iter_rows(min_row=2, values_only=True)}
    assert info_map["Search query"] == "dental clinic"
    assert info_map["Location"] == "Dharampeth, Nagpur"
    assert info_map["Search radius"] == "3 km"
    assert info_map["Total result count"] == 3


def test_sanitize_and_filename():
    assert sanitize_filename_part("Dental   Clinic!!!") == "dental_clinic"
    assert sanitize_filename_part("  ") == "search"
    name = build_filename("Dental Clinic", "Dharampeth, Nagpur")
    assert name.endswith(".xlsx")
    assert "dental_clinic" in name
    assert "dharampeth_nagpur" in name
